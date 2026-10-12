"""Office/HTML adapters. No macros, formula evaluation, OCR or external fetches."""
import json
import zipfile

SUPPORTED={'.pdf','.md','.txt','.csv','.docx','.html','.htm','.pptx','.xlsx','.xls'}


def archive_check(path):
    with zipfile.ZipFile(path) as archive:
        entries=archive.infolist()
        if len(entries)>5000 or sum(i.file_size for i in entries)>100*1024*1024:
            raise ValueError('archive_expansion_limit')
        if any('vbaproject' in i.filename.lower() for i in entries):
            raise ValueError('macros_rejected')


def parse_extra(path,suffix):
    blocks,warnings=[],[]
    headings=[]
    def add(text,location,kind='paragraph',cells=None):
        if text.strip():
            b={'text':text,'normalized':text.replace('\r\n','\n'),'location':location,'headings':list(headings),'kind':kind}
            if cells is not None: b['table_cells']=cells
            blocks.append(b)
        if len(blocks)>20000 or sum(len(b['text']) for b in blocks)>2_000_000:
            raise ValueError('extracted_text_limit')
    def table(rows,location,has_header=False):
        if not rows: return
        width=max(map(len,rows))
        if width>200 or len(rows)>20000: raise ValueError('table_dimensions_limit')
        header=[str(v) if v is not None else '' for v in rows[0]] if has_header else []
        header += ['']*(width-len(header))
        for n,row in enumerate(rows[1:] if has_header else rows,2 if has_header else 1):
            cells=[{'column':f'{i+1}:{header[i]}','value':'' if v is None else str(v)} for i,v in enumerate(row)]
            text=json.dumps({c['column']:c['value'] for c in cells},ensure_ascii=False) if has_header else ' | '.join(c['value'] for c in cells)
            add(text,
                {**location,'row':n,'header':header},'table',cells)
    if suffix in {'.docx','.pptx','.xlsx'}: archive_check(path)
    if suffix=='.docx':
        from docx import Document
        from docx.text.paragraph import Paragraph
        from docx.table import Table
        doc=Document(path)
        stack=[]
        for n,item in enumerate(doc.iter_inner_content(),1):
            if isinstance(item,Paragraph):
                style=item.style.name if item.style else ''
                if style.startswith('Heading ') and style.split()[-1].isdigit():
                    level=int(style.split()[-1])
                    while stack and stack[-1][0]>=level: stack.pop()
                    stack.append((level,item.text));headings[:]=[v for _,v in stack]
                add(item.text,{'paragraph':n})
            elif isinstance(item,Table):
                from docx.oxml.ns import qn
                props=item.rows[0]._tr.trPr if item.rows else None
                marker=props.find(qn('w:tblHeader')) if props is not None else None
                has_header=marker is not None and marker.get(qn('w:val'),'true') not in ('0','false','off')
                table([[c.text for c in row.cells] for row in item.rows],{'table':n},has_header)
        warnings.append('DOCX 无可靠页码；不含页眉页脚、批注、文本框和图片文字。')
        coverage={'body_items':n if 'n' in locals() else 0}
    elif suffix=='.pptx':
        from pptx import Presentation
        deck=Presentation(path)
        if len(deck.slides)>300: raise ValueError('too_many_slides')
        for n,slide in enumerate(deck.slides,1):
            headings[:]=[slide.shapes.title.text] if slide.shapes.title else []
            for i,shape in enumerate(slide.shapes,1):
                if shape.has_table: table([[c.text for c in row.cells] for row in shape.table.rows],{'slide':n,'table':i})
                elif shape.has_text_frame: add(shape.text,{'slide':n,'shape':i})
        warnings.append('PPTX 不含备注、组合形状、SmartArt、嵌入对象和图片文字。')
        coverage={'slides':len(deck.slides)}
    elif suffix in {'.xlsx','.xls'}:
        if suffix=='.xlsx':
            import openpyxl
            book=openpyxl.load_workbook(path,read_only=True,data_only=True,keep_links=False)
            try:
                if len(book.worksheets)>30: raise ValueError('too_many_sheets')
                for sheet in book.worksheets:
                    if sheet.max_row>20000 or sheet.max_column>200: raise ValueError('table_dimensions_limit')
                    headings[:]=[sheet.title]
                    table(list(sheet.iter_rows(values_only=True)),{'sheet':sheet.title},True)
                coverage={'sheets':len(book.worksheets)}
            finally: book.close()
        else:
            import xlrd
            book=xlrd.open_workbook(path,on_demand=True)
            try:
                if book.nsheets>30: raise ValueError('too_many_sheets')
                for sheet in book.sheets():
                    if sheet.nrows>20000 or sheet.ncols>200: raise ValueError('table_dimensions_limit')
                    headings[:]=[sheet.name]
                    table([sheet.row_values(i) for i in range(sheet.nrows)],{'sheet':sheet.name},True)
                coverage={'sheets':book.nsheets}
            finally: book.release_resources()
        warnings.append('首行作为表头；公式只读缓存值，缺缓存为空，不执行公式/宏；复杂合并表头需人工核对。XLS 日期可能是序列数。')
    elif suffix in {'.html','.htm'}:
        from bs4 import BeautifulSoup, NavigableString, Comment, Tag
        from charset_normalizer import from_bytes
        raw=path.read_bytes();decoded=from_bytes(raw).best()
        if decoded is None: raise ValueError('unknown_text_encoding_use_UTF8')
        soup=BeautifulSoup(str(decoded),'html.parser')
        for tag in soup(['script','style','nav','header','footer','aside','iframe','object','noscript','svg']): tag.decompose()
        root=soup.find('main') or soup.find('article') or soup.body or soup
        stack=[]; pending=[]; elements=0
        def flush():
            nonlocal elements
            text=' '.join(''.join(pending).split());pending.clear()
            if text:
                elements+=1;add(text,{'element':elements})
        def html_table(tag):
            nonlocal elements
            rows=[r for r in tag.find_all('tr') if r.find_parent('table') is tag]
            if len(rows)>20000: raise ValueError('table_dimensions_limit')
            grid={}; widths={}; output=[]
            for ri,row in enumerate(rows):
                ci=0
                for cell in row.find_all(['th','td'],recursive=False):
                    while (ri,ci) in grid: ci+=1
                    try: rs=int(cell.get('rowspan',1));cs=int(cell.get('colspan',1))
                    except (ValueError,TypeError): raise ValueError('invalid_table_span')
                    if rs==0: rs=len(rows)-ri
                    if rs<1 or cs<1 or ri+rs>20000 or ci+cs>200: raise ValueError('table_dimensions_limit')
                    value=cell.get_text(' ',strip=True)
                    for rr in range(ri,min(ri+rs,len(rows))):
                        for cc in range(ci,ci+cs):
                            if (rr,cc) in grid: raise ValueError('overlapping_table_span')
                            grid[rr,cc]=value
                            widths[rr]=max(widths.get(rr,0),cc+1)
                    ci+=cs
                width=widths.pop(ri,0)
                output.append([grid.get((ri,col),'') for col in range(width)])
                for col in range(width): grid.pop((ri,col),None)
            first_cells=rows[0].find_all(['th','td'],recursive=False) if rows else []
            header=bool(first_cells) and all(c.name=='th' for c in first_cells) and len(rows)>1
            elements+=1;table(output,{'element':elements},header)
        boundaries={'div','p','section','article','main','li','ul','ol','blockquote','dl','dt','dd','figure','figcaption','address','hr'}
        def walk(node):
            nonlocal elements
            if isinstance(node,Comment): return
            if isinstance(node,NavigableString): pending.append(str(node));return
            if not isinstance(node,Tag): return
            if node.name in ('h1','h2','h3','h4','h5','h6'):
                flush();text=node.get_text('',strip=False).strip();level=int(node.name[1])
                while stack and stack[-1][0]>=level: stack.pop()
                stack.append((level,text));headings[:]=[v for _,v in stack]
                elements+=1;add(text,{'element':elements});return
            if node.name=='table': flush();html_table(node);return
            if node.name=='pre':
                flush();elements+=1;add(node.get_text(),{'element':elements});return
            if node.name=='br': pending.append('\n');return
            if node.name in boundaries: flush()
            for child in node.children: walk(child)
            if node.name in boundaries: flush()
        walk(root);flush()
        warnings.append('HTML 不执行脚本、不获取外链资源；保留正文容器文本和表格网格，覆盖统计仅针对已解析内容。')
        coverage={'elements':len(blocks)}
    else: raise ValueError('unsupported_format')
    if not blocks: raise ValueError('no_text')
    return {'blocks':blocks,'warnings':warnings,'coverage':coverage}
