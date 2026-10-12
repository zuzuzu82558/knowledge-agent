import time
import zipfile

import pytest
from fastapi.testclient import TestClient

from app.config import Settings
from app.main import create_app
from app.parsing import parse,chunk


def write_sample(path,suffix,long=False):
    value='收入120亿元。'+('说明'*1100 if long else '')
    if suffix=='.docx':
        from docx import Document
        doc=Document();doc.add_heading('财务',level=1);doc.add_paragraph(value)
        t=doc.add_table(rows=2,cols=2)
        for c,v in zip([t.cell(0,0),t.cell(0,1),t.cell(1,0),t.cell(1,1)],['公司','指标','星河',value]):c.text=v
        doc.save(path)
    elif suffix=='.pptx':
        from pptx import Presentation
        deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text='财务';slide.placeholders[1].text=value;deck.save(path)
    elif suffix=='.xlsx':
        from openpyxl import Workbook
        book=Workbook();sheet=book.active;sheet.title='财务';sheet.append(['公司','指标']);sheet.append(['星河',value]);book.save(path)
    elif suffix=='.xls':
        import xlwt
        book=xlwt.Workbook();sheet=book.add_sheet('财务')
        for r,row in enumerate([['公司','指标'],['星河',value]]):
            for c,v in enumerate(row):sheet.write(r,c,v)
        book.save(str(path))
    else:
        path.write_text('<html><nav>噪声</nav><main><h1>财务</h1><script>恶意代码</script><p>'+value+'</p></main></html>',encoding='utf-8')


@pytest.mark.parametrize('suffix',['.docx','.pptx','.xlsx','.xls','.html'])
def test_format_normal_long_and_ingest(tmp_path,suffix):
    path=tmp_path/('source'+suffix);write_sample(path,suffix,True)
    parsed=parse(path,suffix);chunks=chunk(parsed)
    assert '120亿元' in ''.join(c['text'] for c in chunks)
    assert all(len(c['text'])<=900 for c in chunks)
    assert all(c['locations'] for c in chunks)
    assert '恶意代码' not in str(parsed) and '噪声' not in str(parsed)
    data=tmp_path/'data'
    with TestClient(create_app(Settings(_env_file=None,data_dir=data))) as client:
        response=client.post('/api/documents',files={'file':(path.name,path.read_bytes())})
        assert response.status_code==202
        rec=response.json()
        for _ in range(300):
            job=client.get('/api/jobs/'+rec['job_id']).json()
            if job['status'] in {'completed','failed'}:break
            time.sleep(.05)
        assert job['status']=='completed',job
        result=client.post('/api/query',json={'question':'收入120亿元','document_ids':[rec['document_id']]}).json()
        assert result['status']=='offline_extracts' and result['evidence']
        for e in result['evidence']:
            source=client.get('/api/evidence/'+e['id']).json()
            assert source['locations']==e['locations']


@pytest.mark.parametrize('suffix',['.docx','.pptx','.xlsx','.xls'])
def test_corrupt_office_rejected(tmp_path,suffix):
    p=tmp_path/('bad'+suffix);p.write_bytes(b'not office')
    with pytest.raises(Exception):parse(p,suffix)


def test_archive_macro_and_formula_cache(tmp_path):
    p=tmp_path/'macro.docx'
    with zipfile.ZipFile(p,'w') as z:z.writestr('word/vbaProject.bin',b'x')
    with pytest.raises(ValueError,match='macros_rejected'):parse(p,'.docx')
    from openpyxl import Workbook
    p=tmp_path/'formula.xlsx';b=Workbook();s=b.active;s.append(['name','formula']);s.append(['星河','=1+1']);b.save(p)
    result=parse(p,'.xlsx')
    assert result['blocks'][0]['table_cells'][1]['value']==''
    assert '公式只读缓存值' in result['warnings'][0]
    p=tmp_path/'bad.html';p.write_text('<script>ignored()</script>',encoding='utf-8')
    with pytest.raises(ValueError,match='no_text'):parse(p,'.html')
