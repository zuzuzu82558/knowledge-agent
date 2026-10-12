"""Acceptance regressions, migrated from the independent review without changing its artifacts.
All model responses below are synthetic fault fixtures, never quality evidence.
"""
import asyncio
import json
import socket
from pathlib import Path
from types import SimpleNamespace

import pytest
from fastapi import HTTPException
from app.config import Settings
from app.service import Service
from app.models import CloudModels
from app.parsing import parse, chunk
from app.evaluation import evidence_metrics

@pytest.fixture(autouse=True)
def offline_network(monkeypatch):
    original=socket.getaddrinfo
    def resolve(host,*a,**k):
        assert host in ('localhost','127.0.0.1','::1',None), 'external network forbidden'
        return original(host,*a,**k)
    connect=socket.socket.connect
    def guarded(sock,address):
        assert not isinstance(address,tuple) or address[0] in ('localhost','127.0.0.1','::1')
        return connect(sock,address)
    monkeypatch.setattr(socket,'getaddrinfo',resolve)
    monkeypatch.setattr(socket.socket,'connect',guarded)

async def ingest(s,text='星河生产芯片。星河收入120亿元。',suffix='.txt'):
    rec=await s.upload('audit'+suffix,suffix,text.encode())
    await asyncio.gather(*s.tasks)
    assert s.store.one('SELECT status FROM jobs WHERE id=?',(rec['job_id'],))['status']=='completed'
    return rec

def fake_models(monkeypatch):
    async def embed(self,texts): return [[1.,.2] for _ in texts]
    async def answer(self,question,evidence):
        return {'status':'answered','claims':[{'text':'fixture claim','citations':[evidence[0]['id']]}]}
    async def structured(self,instruction,data,model=None):
        if 'facts 数组' in instruction:
            return {'facts':[{'subject':{'name':'星河','kind':'company'},'object':{'name':'芯片','kind':'product'},'relation':'produces','quote':'星河生产芯片','assertion':'historical'}]}
        return {'core_viewpoints':[],'risk_points':[],'key_data':[]}
    monkeypatch.setattr(CloudModels,'embed',embed)
    monkeypatch.setattr(CloudModels,'answer',answer)
    monkeypatch.setattr(CloudModels,'structured',structured)

def local(tmp_path):
    return Settings(_env_file=None,data_dir=tmp_path,app_mode='local',model_api_key='synthetic-test-only',embedding_dimension=2)

def test_html_mixed_semantic_and_plain_blocks_preserve_all_text(tmp_path):
    p=tmp_path/'mixed.html'
    p.write_text('<html><meta charset="utf-8"><body><main><h1>财务</h1><div>独有应收款为875万元</div><p>审计意见无保留</p></main></body></html>',encoding='utf-8')
    result=parse(p,'.html')
    assert '独有应收款为875万元' in ''.join(b['text'] for b in result['blocks']), result

@pytest.mark.parametrize('suffix',['.docx','.html','.pptx'])
def test_headerless_key_value_table_keeps_first_row(tmp_path,suffix):
    p=tmp_path/('table'+suffix)
    rows=[['公司','验收星河'],['收入','875万元']]
    if suffix=='.docx':
        from docx import Document
        d=Document();t=d.add_table(rows=2,cols=2)
        for i,row in enumerate(rows):
            for j,v in enumerate(row):t.cell(i,j).text=v
        d.save(p)
    elif suffix=='.pptx':
        from pptx import Presentation
        from pptx.util import Inches
        d=Presentation();s=d.slides.add_slide(d.slide_layouts[6]);t=s.shapes.add_table(2,2,Inches(1),Inches(1),Inches(8),Inches(2)).table
        for i,row in enumerate(rows):
            for j,v in enumerate(row):t.cell(i,j).text=v
        d.save(p)
    else:
        p.write_text('<html><meta charset="utf-8"><table><tr><td>公司</td><td>验收星河</td></tr><tr><td>收入</td><td>875万元</td></tr></table></html>',encoding='utf-8')
    result=parse(p,suffix)
    # The first-row values may survive as invented column names; validate pairing.
    assert any('公司' in b['text'] and '验收星河' in b['text'] and ('公司=验收星河' in b['text'] or '公司 | 验收星河' in b['text'] or '"公司": "验收星河"' in b['text']) for b in result['blocks']), result

def test_html_colspan_does_not_drop_trailing_cells(tmp_path):
    p=tmp_path/'span.html'
    p.write_text('<html><meta charset="utf-8"><table><tr><th colspan="2">年度指标</th></tr><tr><td>2025</td><td>收入875万元</td></tr></table></html>',encoding='utf-8')
    result=parse(p,'.html')
    assert '收入875万元' in ''.join(b['text'] for b in result['blocks']), result

def test_failed_same_config_rebuild_preserves_current_availability(tmp_path,monkeypatch):
    async def run():
        s=Service(Settings(_env_file=None,data_dir=tmp_path,app_mode='offline'))
        try:
            r=await ingest(s);old=await s.query('收入',[r['document_id']]);original=s.index.upsert
            def broken(*a): raise ConnectionError('injected_index_outage')
            monkeypatch.setattr(s.index,'upsert',broken)
            job=await s.rebuild(r['document_id']);await asyncio.gather(*s.tasks)
            assert s.store.one('SELECT status FROM jobs WHERE id=?',(job['job_id'],))['status']=='failed'
            assert s.evidence(old['evidence'][0]['id'])['text']
            try: result=await s.query('收入',[r['document_id']])
            except HTTPException as e: result={'http_status':e.status_code,'detail':e.detail}
            assert result.get('status')=='offline_extracts', result
        finally: await s.close()
    asyncio.run(run())

def test_failed_graph_rebuild_invalidates_previous_graph_answer(tmp_path,monkeypatch):
    fake_models(monkeypatch)
    async def run():
        s=Service(local(tmp_path))
        try:
            r=await ingest(s);await s.build_graph(r['document_id'])
            first=await s.query('星河芯片关系',[r['document_id']],route='relation',strategy='graph')
            assert first.get('graph_paths') and not first['warnings']
            def broken(*a): raise ConnectionError('injected_graph_storage_failure')
            monkeypatch.setattr(s.graph,'replace',broken)
            failed=await s.build_graph(r['document_id']);assert failed['status']=='failed'
            second=await s.query('星河芯片关系',[r['document_id']],route='relation',strategy='graph')
            assert second['cache']['status']=='miss' and 'graph_not_complete_hybrid_fallback' in second['warnings'], {'cache':second['cache'],'warnings':second['warnings'],'build':failed}
        finally: await s.close()
    asyncio.run(run())

def test_failed_aggregation_does_not_publish_ready_summary_status(tmp_path,monkeypatch):
    fake_models(monkeypatch)
    async def invalid(self,*a,**k): return {'invalid':'aggregate'}
    monkeypatch.setattr(CloudModels,'structured',invalid)
    async def run():
        s=Service(local(tmp_path))
        try:
            r=await ingest(s);summary=await s.query('全文洞察',[r['document_id']],route='summary')
            assert summary['status']=='partial_summary'
            status=s.store.one('SELECT summary_status FROM generations WHERE id=?',(r['generation_id'],))['summary_status']
            assert status=='partial', {'api':summary['status'],'aggregate':summary['insights']['status'],'stored':status}
        finally: await s.close()
    asyncio.run(run())

def test_fallback_answer_not_served_as_normal_primary_cache(tmp_path,monkeypatch):
    import litellm
    class RateLimitError(Exception): pass
    async def embedding(self,texts):return [[1.,.2] for _ in texts]
    async def completion(model,messages,**kwargs):
        if model=='primary':raise RateLimitError('synthetic')
        evidence=json.loads(messages[1]['content'])['evidence']
        return SimpleNamespace(model='backup',usage=None,choices=[SimpleNamespace(message=SimpleNamespace(content=json.dumps({'refused':False,'claims':[{'text':'fixture','citations':[evidence[0]['id']]}]})))])
    monkeypatch.setattr(CloudModels,'embed',embedding);monkeypatch.setattr(litellm,'acompletion',completion)
    async def run():
        s=Service(local(tmp_path).model_copy(update={'llm_model':'primary','fallback_model':'backup'}))
        try:
            r=await ingest(s);first=await s.query('收入',[r['document_id']]);assert first['usage'][0]['model']=='backup'
            second=await s.query('收入',[r['document_id']])
            assert second['cache']['status']!='exact', {'cache':second['cache'],'warnings':second['warnings'],'policy':second['policy'],'calls':second['model_calls']}
        finally:await s.close()
    asyncio.run(run())

def test_contextual_golden_anchors_reject_wrong_subject():
    from pathlib import Path
    rows=[json.loads(line) for line in Path('eval/golden.jsonl').read_text(encoding='utf-8').splitlines()]
    question=next(row for row in rows if row['id']=='g01')
    wrong=[{'text':'竞争对手总部位于苏州，星河总部未披露。'}]
    score=evidence_metrics(question['evidence_groups'],wrong,wrong)
    assert score['all_required_anchors_present'] is False
    assert score['semantic_support']=='not_judged'
    # A raw city token still matches lexically, but no longer masquerades as
    # a semantic all-necessary-evidence assertion.
    lexical=evidence_metrics([['苏州']],wrong,wrong)
    assert lexical['all_required_anchors_present'] is True
    assert 'all_necessary_evidence' not in lexical
    assert lexical['semantic_support']=='not_judged'


def test_scope_cache_revision_and_history_survive_restart(tmp_path):
    async def run():
        settings=Settings(_env_file=None,data_dir=tmp_path,app_mode='offline')
        s=Service(settings)
        a=await ingest(s);b=await ingest(s,'另一公司收入999亿元。')
        first=await s.query('收入',[a['document_id']]);second=await s.query('收入',[a['document_id']])
        assert second['cache']['status']=='exact'
        assert all(e['document']==a['document_id'] for e in second['evidence'])
        up=await s.upload('audit.txt','.txt','星河收入150亿元。'.encode(),a['document_id']);await asyncio.gather(*s.tasks)
        current=await s.query('收入',[a['document_id']]);assert current['cache']['status']=='miss'
        assert s.evidence(first['evidence'][0]['id'])['text']==first['evidence'][0]['text']
        await s.close();s=Service(settings)
        try:
            assert s.run(first['run_id'])['snapshot']==first['snapshot']
            assert (await s.query('收入',[a['document_id']]))['snapshot']['versions']==[up['version_id']]
            await s.delete(a['document_id'])
            with pytest.raises(HTTPException):s.run(first['run_id'])
            assert (await s.query('收入',[b['document_id']]))['evidence']
        finally:await s.close()
    asyncio.run(run())

def test_csv_multiline_and_markdown_fences_preserve_locations(tmp_path):
    p=tmp_path/'multi.csv';p.write_text('公司,备注,金额\n"星河,一部","首行\n续行",-12.50\n',encoding='utf-8')
    b=parse(p,'.csv')['blocks'][0];assert b['location']['line_end']==3 and b['table_cells'][2]['value']=='-12.50'
    p=tmp_path/'fence.md';p.write_text('# 总览\n### A\n```text\n# 伪标题\n```\n### B\n唯一事实B\n',encoding='utf-8')
    b=next(b for b in parse(p,'.md')['blocks'] if '唯一事实B' in b['text']);assert b['headings']==['总览','B']

def test_real_pdf_table_and_mixed_scan_coverage(tmp_path):
    # Real PDF drawing primitives, no extra generator dependency or parser mock.
    p=tmp_path/'table.pdf'
    stream=b'BT /F1 12 Tf 50 800 Td (Audit financial table for year 2025) Tj ET\n'
    for y in (740,710,680):stream+=f'50 {y} m 350 {y} l S\n'.encode()
    for x in (50,200,350):stream+=f'{x} 680 m {x} 740 l S\n'.encode()
    for x,y,text in [(60,720,'Company'),(210,720,'Revenue'),(60,690,'AuditStar'),(210,690,'-12.50 USD')]:
        stream+=f'BT /F1 12 Tf {x} {y} Td ({text}) Tj ET\n'.encode()
    empty=b'BT /F1 12 Tf 50 800 Td (x) Tj ET'
    objects=[b'<< /Type /Catalog /Pages 2 0 R >>',b'<< /Type /Pages /Kids [3 0 R 6 0 R] /Count 2 >>',
        b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 600 850] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>',
        b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
        b'<< /Length '+str(len(stream)).encode()+b' >>\nstream\n'+stream+b'\nendstream',
        b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 600 850] /Resources << /Font << /F1 4 0 R >> >> /Contents 7 0 R >>',
        b'<< /Length '+str(len(empty)).encode()+b' >>\nstream\n'+empty+b'\nendstream']
    data=b'%PDF-1.4\n';offsets=[]
    for i,obj in enumerate(objects,1):offsets.append(len(data));data+=f'{i} 0 obj\n'.encode()+obj+b'\nendobj\n'
    xref=len(data);data+=b'xref\n0 8\n0000000000 65535 f \n'+b''.join(f'{v:010d} 00000 n \n'.encode() for v in offsets)
    data+=f'trailer\n<< /Size 8 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF'.encode();p.write_bytes(data)
    result=parse(p,'.pdf');assert result['coverage']=={'pages':2,'text_pages':1}
    assert result['warnings'] and any(b['kind']=='table' and '-12.50 USD' in b['text'] for b in result['blocks'])

def test_interrupted_job_restart_requires_explicit_retry(tmp_path):
    async def run():
        settings=Settings(_env_file=None,data_dir=tmp_path,app_mode='offline');s=Service(settings)
        r=await ingest(s)
        s.store.execute("UPDATE jobs SET status='running',stage='parse' WHERE id=?",(r['job_id'],))
        await s.close();s=Service(settings)
        try:
            row=s.store.one('SELECT * FROM jobs WHERE id=?',(r['job_id'],));assert row['status']=='interrupted'
            await s.retry(r['job_id']);await asyncio.gather(*s.tasks)
            assert s.store.one('SELECT status FROM jobs WHERE id=?',(r['job_id'],))['status']=='completed'
        finally:await s.close()
    asyncio.run(run())
